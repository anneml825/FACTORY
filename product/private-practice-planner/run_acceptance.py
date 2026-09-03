"""Runs the pre-registered acceptance tests against the built workbook."""
from build import build, PR, CO, RE
from evaluate import Eval
import tests

# Named so a layout change is a one-line edit, not a search-and-replace.
SLOTS_WK, ATTENDED, REVENUE = "B20", "B21", "B22"
AVG_HELD, AVG_SLOT, TIERWARN = "B23", "B24", "B26"
COSTS_TOTAL = "B32"
TAKE_HOME, DIFF, NEEDED = "B8", "B10", "B13"
TO_COVER, TO_TARGET, SLOTS_NEEDED, TAXWARN = "B14", "B15", "B16", "B17"
REDUCED, REDWARN = "B24", "B25"

results = []


def ev(o, name):
    path = f"dist/tests/{name}.xlsx"
    build(path, o)
    return Eval(path)


def check(name, desc, ok, detail):
    results.append((name, "PASS" if ok else "FAIL", desc, detail))


def close(a, b, tol=1e-6):
    return a not in ("", None) and abs(float(a) - b) <= tol


# T1 — revenue, attendance and the two averages agree with the model
e, m = ev(None, "T1"), tests.model()
check("T1", "revenue, sessions and both averages match the independent model",
      close(e.value(PR, REVENUE), m["revenue"]) and close(e.value(PR, ATTENDED), m["attended"])
      and close(e.value(PR, AVG_HELD), m["avg_held"]) and close(e.value(PR, AVG_SLOT), m["avg_slot"]),
      f"revenue {e.value(PR, REVENUE)} vs {m['revenue']}, per-slot {e.value(PR, AVG_SLOT)}")

# T2 — the per-slot average is strictly below the per-session-held average
check("T2", "a slot is worth less than a session held, because of no-shows",
      float(e.value(PR, AVG_SLOT)) < float(e.value(PR, AVG_HELD)),
      f"slot {e.value(PR, AVG_SLOT)} < held {e.value(PR, AVG_HELD)}")

# T3 — 100% attendance collapses the two averages together
e3, m3 = ev({"tiers": [("Full fee", 150, 10, 1.0), ("", "", "", ""), ("", "", "", ""),
                       ("", "", "", ""), ("", "", "", ""), ("", "", "", "")]}, "T3"), \
         tests.model({"tiers": [("Full fee", 150, 10, 1.0), ("", "", "", ""), ("", "", "", ""),
                                ("", "", "", ""), ("", "", "", ""), ("", "", "", "")]})
check("T3", "at 100% attendance the two averages are equal and both equal the fee",
      close(e3.value(PR, AVG_SLOT), 150) and close(e3.value(PR, AVG_HELD), 150)
      and close(e3.value(PR, AVG_SLOT), m3["avg_slot"]),
      f"{e3.value(PR, AVG_SLOT)} / {e3.value(PR, AVG_HELD)}")

# T4 — costs roll up monthly*12 + annual
check("T4", "annual costs are monthly x 12 plus the annual items",
      close(e.value(CO, COSTS_TOTAL), m["costs"]),
      f"{e.value(CO, COSTS_TOTAL)} vs {m['costs']}")

# T5 — take-home and the gap to target
check("T5", "take-home and difference from target match the model",
      close(e.value(RE, TAKE_HOME), m["take_home"]) and close(e.value(RE, DIFF), m["diff"]),
      f"take-home {e.value(RE, TAKE_HOME)} vs {m['take_home']}")

# T6 — revenue needed inverts back to exactly the target take-home
needed = float(e.value(RE, NEEDED))
implied = (needed - m["costs"]) * (1 - 0.25)
check("T6", "revenue needed, run back through costs and tax, lands on the target",
      close(needed, m["needed"]) and abs(implied - 60000) < 1e-6,
      f"needed {needed:.2f} implies take-home {implied:.2f}")

# T7 — session counts round UP, never down
check("T7", "sessions to cover costs and to reach target both round up",
      close(e.value(RE, TO_COVER), m["to_cover"]) and close(e.value(RE, TO_TARGET), m["to_target"])
      and close(e.value(RE, SLOTS_NEEDED), m["slots_needed"]),
      f"cover {e.value(RE, TO_COVER)}, target {e.value(RE, TO_TARGET)}, slots/wk {e.value(RE, SLOTS_NEEDED)}")

# T8 — the sliding-scale answer is correct AND self-consistent when fed back
red = float(e.value(RE, REDUCED))
full_n = 22 - red
rev_at = 46 * 0.90 * (full_n * 150 + red * 80)
check("T8", "reduced-fee slots match the model and still clear the revenue needed",
      close(e.value(RE, REDUCED), m["reduced"]) and rev_at >= needed - 1e-6,
      f"{red:.0f} reduced slots -> revenue {rev_at:.2f} vs needed {needed:.2f}")

# T9 — one more reduced slot would break the target. The answer is the true maximum.
rev_one_more = 46 * 0.90 * ((full_n - 1) * 150 + (red + 1) * 80)
check("T9", "one further reduced slot would miss the target, so the answer is maximal",
      rev_one_more < needed,
      f"one more gives {rev_one_more:.2f}, below needed {needed:.2f}")

# T10 — reduced fee at or above full fee is refused, with no number
e10 = ev({"red_fee": 150}, "T10")
check("T10", "a reduced fee equal to the full fee is refused, not computed",
      e10.value(RE, REDUCED) == "" and "below the full fee" in str(e10.value(RE, REDWARN)),
      f"value {e10.value(RE, REDUCED)!r}, warning {str(e10.value(RE, REDWARN))[:44]!r}")

# T11 — an unreachable target returns zero slots and says so
e11 = ev({"target": 400000}, "T11")
check("T11", "an unreachable target yields zero reduced slots and an explanation",
      close(e11.value(RE, REDUCED), 0)
      and "full fee" in str(e11.value(RE, REDWARN)),
      f"{e11.value(RE, REDUCED)!r} / {str(e11.value(RE, REDWARN))[:44]!r}")

# T12 — a target the full caseload covers easily caps at the slot count
e12 = ev({"target": 1000}, "T12")
check("T12", "an easily covered target caps at the number of slots, and says so",
      close(e12.value(RE, REDUCED), 22) and "every slot" in str(e12.value(RE, REDWARN)),
      f"{e12.value(RE, REDUCED)!r} / {str(e12.value(RE, REDWARN))[:44]!r}")

# T13 — empty caseload produces no errors anywhere and warns
blank = [("", "", "", "")] * 6
e13 = ev({"tiers": blank}, "T13")
vals13 = [e13.value(s, f"B{r}") for s in (PR,) for r in range(20, 27)]
check("T13", "an empty caseload produces no error values and prompts for input",
      not any("#" in str(v) for v in vals13) and "at least one" in str(e13.value(PR, TIERWARN)),
      f"{[str(v)[:12] for v in vals13]}")

# T14 — a tier missing its attendance is caught by the row check
bad = [("Full fee", 150, 10, ""), ("", "", "", ""), ("", "", "", ""),
       ("", "", "", ""), ("", "", "", ""), ("", "", "", "")]
e14 = ev({"tiers": bad}, "T14")
check("T14", "a tier missing its attendance rate is flagged on that row",
      "attendance rate" in str(e14.value(PR, "G12")).lower(),
      f"{str(e14.value(PR, 'G12'))[:48]!r}")

# T15 — attendance above 100% is refused on the row
over = [("Full fee", 150, 10, 1.4), ("", "", "", ""), ("", "", "", ""),
        ("", "", "", ""), ("", "", "", ""), ("", "", "", "")]
e15 = ev({"tiers": over}, "T15")
check("T15", "an attendance rate above 100% is flagged",
      "above 100%" in str(e15.value(PR, "G12")),
      f"{str(e15.value(PR, 'G12'))[:48]!r}")

# --- full-book scan -----------------------------------------------------------
import openpyxl, re
wb = openpyxl.load_workbook("dist/private-practice-planner.xlsx")
formulas = [(ws.title, c.coordinate, c.value) for ws in wb for r in ws.iter_rows()
            for c in r if isinstance(c.value, str) and c.value.startswith("=")]
broken = [f"{s}!{co}" for s, co, v in formulas
          if v.count("(") != v.count(")") or v.count('"') % 2]

print()
for name, status, desc, detail in results:
    print(f"{name:5} {status}  {desc}")
    print(f"        {detail}")
print()
print(f"tabs: {wb.sheetnames}")
print(f"formulas: {len(formulas)}   bracket/quote problems: {len(broken) or 'none'}")

agree = all(s == "PASS" for _, s, _, _ in results)
print(f"\nindependent-model cross-check across all cases: {'AGREE' if agree else 'DISAGREE'}")
passed = sum(1 for _, s, _, _ in results if s == "PASS")
print(f"{passed}/{len(results)} acceptance tests pass")
