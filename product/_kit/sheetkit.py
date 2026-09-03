"""
Shared machinery for Factory workbooks.

Extracted from the first three products rather than designed up front, so it
carries the decisions that were actually paid for:

  - fullCalcOnLoad, because a shipped workbook once rendered one stale value
    beside a correct one
  - a warning row is the only place a full sentence belongs in a spreadsheet
  - every default visible and editable, colour-coded by kind
  - blank inputs must never produce #DIV/0! or #VALUE!

Copy rule for anything built on this: labels are labels. A few words, no
sentences, no commentary about how the sheet works. Explanation lives in the
Start Here PDF.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

# --- palette ------------------------------------------------------------------
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")   # type here
CALC_FILL = PatternFill("solid", fgColor="F1F3F4")    # calculated
ASSUM_FILL = PatternFill("solid", fgColor="E3F2FD")   # editable assumption


def head_fill(hex_):
    return PatternFill("solid", fgColor=hex_)


# --- type ---------------------------------------------------------------------
def fonts(accent="1B3A4B"):
    return dict(
        H1=Font(bold=True, size=14, color=accent),
        H2=Font(bold=True, size=10, color="FFFFFF"),
        BOLD=Font(bold=True),
        NOTE=Font(size=9, color="607D8B", italic=True),
        WARN=Font(bold=True, size=10, color="C62828"),
        BIG=Font(bold=True, size=12, color=accent),
    )


THIN = Border(*[Side("thin", color="CFD8DC")] * 4)

MONEY = '#,##0.00'
PCT = '0%'
PCT1 = '0.0%'
NUM = '0'
NUM1 = '0.0'
NUM2 = '0.00'
DATE = 'yyyy-mm-dd'


class Sheet:
    """A thin wrapper that keeps the per-product builders short and uniform."""

    def __init__(self, wb, title, accent="1B3A4B", widths=None):
        self.ws = wb.create_sheet(title)
        self.f = fonts(accent)
        self.accent = accent
        self.ws["A1"] = title
        self.ws["A1"].font = self.f["H1"]
        for col, w in (widths or {}).items():
            self.ws.column_dimensions[col].width = w

    def title(self, text):
        self.ws["A1"] = text
        self.ws["A1"].font = self.f["H1"]
        return self

    def section(self, row, text, span=4):
        for i in range(span):
            self.ws.cell(row=row, column=1 + i).fill = head_fill(self.accent)
        c = self.ws.cell(row=row, column=1, value=text)
        c.font = self.f["H2"]
        return self

    def label(self, row, text, note=None, col=1, note_col=4, bold=False):
        c = self.ws.cell(row=row, column=col, value=text)
        if bold:
            c.font = self.f["BOLD"]
        if note:
            self.ws.cell(row=row, column=note_col, value=note).font = self.f["NOTE"]
        return self

    def headers(self, row, pairs):
        for col, txt in pairs:
            self.ws.cell(row=row, column=col, value=txt).font = self.f["BOLD"]
        return self

    def inp(self, row, value, fmt=None, col=2, fill=INPUT_FILL):
        c = self.ws.cell(row=row, column=col, value=value)
        c.fill, c.border = fill, THIN
        if fmt:
            c.number_format = fmt
        return c

    def assum(self, row, value, fmt=None, col=2):
        return self.inp(row, value, fmt, col, ASSUM_FILL)

    def calc(self, row, formula, fmt=None, big=False, col=2):
        c = self.ws.cell(row=row, column=col, value=formula)
        c.fill, c.border = CALC_FILL, THIN
        if fmt:
            c.number_format = fmt
        if big:
            c.font = self.f["BIG"]
        return c

    def warn(self, row, formula, col=2):
        c = self.ws.cell(row=row, column=col, value=formula)
        c.font = self.f["WARN"]
        return c

    def note(self, row, text, col=3):
        self.ws.cell(row=row, column=col, value=text)
        return self

    def dropdown(self, ref, options):
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        self.ws.add_data_validation(dv)
        dv.add(ref)
        return self

    def color_key(self, row, span=3):
        """Every product ships the same legend, so it lives here."""
        self.section(row, "Cell colors", span)
        for i, (fill, txt) in enumerate([(INPUT_FILL, "Enter your own numbers"),
                                         (CALC_FILL, "Calculated automatically"),
                                         (ASSUM_FILL, "Editable assumption")]):
            c = self.ws.cell(row=row + 1 + i, column=1)
            c.fill, c.border = fill, THIN
            self.ws.cell(row=row + 1 + i, column=3, value=txt)
        return self


def new_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def finish(wb):
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
    # Recalculate on open rather than showing whatever was cached at build time.
    wb.calculation.fullCalcOnLoad = True
    return wb


def g(o, key, default):
    """Test override lookup. Products take an optional dict so acceptance tests
    can rebuild the same workbook with different inputs."""
    return default if not o or key not in o else o[key]


def blank_guard(cells):
    """OR($A1="",$B1="") for a list of refs — the guard that keeps a blank
    input from producing #DIV/0! or #VALUE!."""
    return "OR(" + ",".join(f'{c}=""' for c in cells) + ")"
